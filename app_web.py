import streamlit as st
import pandas as pd
import numpy as np
import io
import pulp
import traceback
import random
import openpyxl
from datetime import datetime

# ==========================================
# 1. 網頁頁面配置
# ==========================================
st.set_page_config(page_title="段考監考終極自動化", page_icon="🏫", layout="wide")
st.title("🏫 試務組-段考監考全自動化系統 (強制配對版)")
st.info("💡 終極更新：已加裝「檔案真實日期自動萃取器」，徹底解決因忘記調整網頁日期而導致的配對空白問題！")

# --- 初始化狀態 ---
if 'results' not in st.session_state:
    st.session_state['results'] = None
if 'uploader_key' not in st.session_state:
    st.session_state['uploader_key'] = 0

# ==========================================
# 2. 輔助功能定義
# ==========================================
def to_excel_bytes(df, header_df=None):
    output = io.BytesIO()
    if header_df is not None:
        df.columns = header_df.columns
        final_out = pd.concat([header_df, df], ignore_index=True)
    else:
        final_out = df
    final_out = final_out.fillna("")
    with pd.ExcelWriter(output, engine='xlsxwriter') as writer:
        final_out.to_excel(writer, index=False, header=False)
    return output.getvalue()

# 【字串全面淨化器】
def clean_str(s):
    if pd.isna(s) or s is None: return ""
    s = str(s).strip().replace('ㄧ', '一').replace(' ', '').replace('　', '').replace('\n', '').replace('\r', '')
    s = s.translate(str.maketrans('１２３４５６７８９０', '1234567890'))
    return s

def normalize_subject(s):
    s = clean_str(s)
    aliases = {'國文':'國語文', '英文':'英語文', '公社':'公民與社會', '公民':'公民與社會', 
               '地科':'地球科學', '健護':'健康與護理', '護理':'健康與護理', 
               '國防':'全民國防教育', '生科':'生活科技', '應數':'應用數學'}
    return aliases.get(s, s)

def get_teacher_fuzzy(cls, subj, course_dict):
    if (cls, subj) in course_dict: return course_dict[(cls, subj)]
    clean_target = subj.replace('選修', '').replace('彈性學習', '').replace('補強', '').replace('-', '')
    for (c, s), t in course_dict.items():
        if c == cls:
            s_clean = s.replace('選修', '').replace('彈性學習', '').replace('補強', '').replace('-', '')
            if clean_target and (clean_target in s_clean or s_clean in clean_target):
                return t
    if len(clean_target) >= 2:
        for (c, s), t in course_dict.items():
            if c == cls and clean_target[:2] in s:
                return t
    return ""

# ==========================================
# 3. 介面佈局
# ==========================================
st.divider()
col1, col2 = st.columns([1, 1], gap="large")

with col1:
    st.subheader("📂 1. 上傳排考與標籤資料")
    file_quota = st.file_uploader("1️⃣ 監考堂數.xlsx", type=['xlsx'], key=f"f1_{st.session_state['uploader_key']}")
    file_list = st.file_uploader("2️⃣ 監考名單.xlsx", type=['xlsx'], key=f"f2_{st.session_state['uploader_key']}")
    file_type = st.file_uploader("3️⃣ 監考類型總數.xlsx", type=['xlsx'], key=f"f3_{st.session_state['uploader_key']}")
    file_pub = st.file_uploader("4️⃣ 監考總表公布版.xlsx (範本)", type=['xlsx'], key=f"f4_{st.session_state['uploader_key']}")
    file_assign = st.file_uploader("5️⃣ 監考一覽表.xlsx (班級分配範本)", type=['xlsx'], key=f"f5_{st.session_state['uploader_key']}")
    st.write("---")
    file_course = st.file_uploader("6️⃣ 配課表.xlsx (多工作表)", type=['xlsx'], key=f"f6_{st.session_state['uploader_key']}")
    file_label = st.file_uploader("7️⃣ 標籤列印.xlsx (試卷袋範本)", type=['xlsx'], key=f"f7_{st.session_state['uploader_key']}")

with col2:
    st.subheader("⚙️ 2. 考試設定與特許名單")
    selected_sheet = None
    if file_quota:
        xls = pd.ExcelFile(file_quota)
        selected_sheet = st.selectbox("👇 選擇考試項目：", xls.sheet_names)
    
    flex_names = []
    if file_list:
        temp_df = pd.read_excel(file_list, header=None).fillna("")
        teacher_list = temp_df.iloc[2:, 1].astype(str).str.strip().tolist()
        teacher_list = [t for t in teacher_list if t != "" and t != "nan"]
        flex_names = st.multiselect("🛡️ 優先時數不大於名單：", options=teacher_list)

    st.write("")
    c_d1, c_d2 = st.columns(2)
    with c_d1: d1_date = st.date_input("📅 第一天日期(選填)：", datetime.now())
    with c_d2: d2_date = st.date_input("📅 第二天日期(選填)：", datetime.now())
    
    force_run = st.checkbox("⚠️ 忽略健檢警告，強制執行")
    if st.button("🗑️ 清除所有設定", use_container_width=True):
        st.session_state['results'] = None
        st.session_state['uploader_key'] += 1
        st.rerun()

# ==========================================
# 4. 核心演算法執行
# ==========================================
st.divider()

if st.button("🚀 啟動終極全自動排班系統", type="primary", use_container_width=True):
    if not all([file_quota, file_list, file_type, file_assign]):
        st.error("🚨 請至少確認【1, 2, 3, 5】號基礎檔案皆已上傳！")
    else:
        try:
            df_quota = pd.read_excel(file_quota, sheet_name=selected_sheet).fillna("")
            quota_dict = dict(zip(df_quota.iloc[:, 0].astype(str).str.strip(), pd.to_numeric(df_quota.iloc[:, 1], errors='coerce').fillna(0)))
            
            df_type = pd.read_excel(file_type, header=None).fillna("")
            req_matrix = {'△': [0]*10, '※': [0]*10}
            for i in range(2, len(df_type)):
                row_name = str(df_type.iloc[i, 0]).strip()
                if row_name in ['△', '※']:
                    req_matrix[row_name] = pd.to_numeric(df_type.iloc[i, 1:11], errors='coerce').fillna(0).astype(int).tolist()

            # --- 讀取基礎資料與【防呆真實日期萃取】 ---
            df_assign_raw = pd.read_excel(file_assign, header=None).fillna("")
            try:
                # 直接從您上傳的《監考一覽表》讀取真實的考試日期，取代網頁介面的輸入！
                sys_d1_raw = str(df_assign_raw.iloc[0, 1])[:10].replace('/', '-')
                sys_d2_raw = str(df_assign_raw.iloc[0, 6])[:10].replace('/', '-')
            except:
                sys_d1_raw = d1_date.strftime('%Y-%m-%d')
                sys_d2_raw = d2_date.strftime('%Y-%m-%d')

            # 產生報表用的顯示日期 (如 05月14日)
            display_d1 = f"{sys_d1_raw[-5:-3]}月{sys_d1_raw[-2:]}日" if len(sys_d1_raw) >= 5 else d1_date.strftime('%m月%d日')
            display_d2 = f"{sys_d2_raw[-5:-3]}月{sys_d2_raw[-2:]}日" if len(sys_d2_raw) >= 5 else d2_date.strftime('%m月%d日')

            df_list_raw = pd.read_excel(file_list, header=None).fillna("")
            header_df = df_list_raw.iloc[0:2].copy().astype(str).replace('nan', '')
            for c in range(3, 8): header_df.iloc[0, c] = display_d1
            for c in range(8, 13): header_df.iloc[0, c] = display_d2
            
            df_list = df_list_raw.iloc[2:].copy()
            teachers = df_list.iloc[:, 1].astype(str).str.strip().tolist()

            # --- PuLP 運算 ---
            with st.spinner("🧠 正在生成完美監考總表..."):
                prob = pulp.LpProblem("Scheduling", pulp.LpMinimize)
                vX = {}; vY = {}
                for i in range(len(teachers)):
                    vX[i] = {}; vY[i] = {}
                    for j in range(10):
                        vX[i][j] = pulp.LpVariable(f"X_{i}_{j}", cat='Binary')
                        vY[i][j] = pulp.LpVariable(f"Y_{i}_{j}", cat='Binary')
                
                penalty = 0
                for i, t in enumerate(teachers):
                    tgt = int(quota_dict.get(t, 0))
                    act = pulp.lpSum([vX[i][k] + vY[i][k]*2 for k in range(10)])
                    prob += act <= tgt
                    dfct = pulp.LpVariable(f"dfct_{i}", 0)
                    prob += act + dfct == tgt
                    penalty += dfct * (1 if t in flex_names else 1000)
                    for j in range(10):
                        prob += vX[i][j] + vY[i][j] <= 1
                        cell_val = str(df_list.iloc[i, j+3]).strip()
                        if cell_val != "" and cell_val != "nan":
                            prob += vX[i][j] == 0; prob += vY[i][j] == 0
                    prob += vX[i][1] >= vY[i][0]
                    prob += vX[i][6] >= vY[i][5]
                for j in range(10):
                    prob += pulp.lpSum([vX[i][j] for i in range(len(teachers))]) == req_matrix['△'][j]
                    prob += pulp.lpSum([vY[i][j] for i in range(len(teachers))]) == req_matrix['※'][j]
                prob += penalty
                prob.solve()

                schedule_dict = {}
                df_out_master = df_list.copy()
                for i, t in enumerate(teachers):
                    res = []
                    for j in range(10):
                        val = str(df_list.iloc[i, j+3]).strip()
                        if val == "" or val == "nan":
                            if vX[i][j].varValue == 1: val = "△"
                            elif vY[i][j].varValue == 1: val = "※"
                            else: val = "" 
                        res.append(val)
                        df_out_master.iloc[i, j+3] = val
                    schedule_dict[t] = res

            # --- 監考一覽表分配邏輯 ---
            with st.spinner("🎯 執行班級自動分配..."):
                assign_header = df_assign_raw.iloc[0:2].copy().astype(str).replace('nan', '')
                for c in range(1, 6): assign_header.iloc[0, c] = display_d1
                for c in range(6, 11): assign_header.iloc[0, c] = display_d2

                df_assign = df_assign_raw.iloc[2:].copy()
                class_names_raw = df_assign.iloc[:, 0].tolist()
                
                norm_class_names = [clean_str(c) for c in class_names_raw]
                assign_map = {name: idx for idx, name in enumerate(norm_class_names)}
                
                assigned_matrix = np.empty((len(class_names_raw), 10), dtype=object)
                
                for day_start in [0, 5]:
                    j1 = day_start
                    proctors_j1 = [t for t in teachers if schedule_dict[t][j1] in ["△", "※"]]
                    random.shuffle(proctors_j1)
                    for idx, p in zip(range(len(class_names_raw)), proctors_j1): 
                        assigned_matrix[idx, j1] = p
                    
                    j2 = day_start + 1
                    proctors_j2 = [t for t in teachers if schedule_dict[t][j2] in ["△", "※"]]
                    bound = {}
                    for idx in range(len(class_names_raw)):
                        p_prev = assigned_matrix[idx, j1]
                        if p_prev is not None and schedule_dict[p_prev][j1] == "※" and schedule_dict[p_prev][j2] == "△":
                            assigned_matrix[idx, j2] = p_prev
                            bound[p_prev] = True
                    
                    rem = [p for p in proctors_j2 if p not in bound]
                    random.shuffle(rem)
                    r_idx = 0
                    for idx in range(len(class_names_raw)):
                        if assigned_matrix[idx, j2] is None and r_idx < len(rem):
                            assigned_matrix[idx, j2] = rem[r_idx]
                            r_idx += 1

                    for offset in [2, 3, 4]:
                        curr_j = day_start + offset
                        proctors = [t for t in teachers if schedule_dict[t][curr_j] in ["△", "※"]]
                        random.shuffle(proctors)
                        for idx, p in zip(range(len(class_names_raw)), proctors): 
                            assigned_matrix[idx, curr_j] = p

                for r in range(len(class_names_raw)):
                    for c in range(10): df_assign.iloc[r, c+1] = assigned_matrix[r, c]

            # --- 公布版套印 ---
            pub_bytes = None
            if file_pub:
                with st.spinner("🖨️ 正在無縫套印公布版..."):
                    wb = openpyxl.load_workbook(file_pub)
                    ws = wb.active
                    h_row = -1; t_cols = []
                    for r in range(1, min(20, ws.max_row + 1)):
                        for c in range(1, ws.max_column + 1):
                            val = ws.cell(row=r, column=c).value
                            if val and "教師" in str(val):
                                h_row = r; t_cols.append(c)
                        if h_row != -1: break
                    if h_row != -1:
                        for c in t_cols:
                            if h_row - 1 >= 1:
                                ws.cell(row=h_row-1, column=c+2).value = display_d1
                                ws.cell(row=h_row-1, column=c+7).value = display_d2
                            for r in range(h_row+1, ws.max_row + 1):
                                t_val = ws.cell(row=r, column=c).value
                                if t_val:
                                    name = str(t_val).strip()
                                    if name in schedule_dict:
                                        for j in range(5): ws.cell(row=r, column=c+2+j).value = schedule_dict[name][j]
                                        for j in range(5): ws.cell(row=r, column=c+7+j).value = schedule_dict[name][j+5]
                    out_pub = io.BytesIO()
                    wb.save(out_pub)
                    pub_bytes = out_pub.getvalue()

            # --- 標籤列印合成 (採用強制真實日期與絕對座標) ---
            label_bytes = None
            if file_course and file_label:
                with st.spinner("🏷️ 啟動【無敵真實日期座標綁定】，精準合成標籤..."):
                    
                    course_dict = {}
                    xls_course = pd.ExcelFile(file_course)
                    for sheet in xls_course.sheet_names:
                        df_c = pd.read_excel(file_course, sheet_name=sheet, header=None).fillna("")
                        h_idx = 0
                        for r in range(min(5, len(df_c))):
                            row_str = "".join(str(x) for x in df_c.iloc[r, :])
                            if "科目" in row_str or "一1" in row_str or "二1" in row_str or "三1" in row_str:
                                h_idx = r; break
                        classes_in_sheet = [clean_str(x) for x in df_c.iloc[h_idx, :]]
                        
                        for r_idx in range(h_idx + 1, len(df_c)):
                            row = df_c.iloc[r_idx, :]
                            subj_raw = str(row.iloc[0])
                            if not subj_raw: continue
                            subj_norm = normalize_subject(subj_raw)
                            
                            for c_idx in range(1, len(row)):
                                cls_raw = classes_in_sheet[c_idx]
                                teacher = clean_str(row.iloc[c_idx])
                                if teacher and cls_raw:
                                    course_dict[(cls_raw, subj_norm)] = teacher
                    
                    wb_label = openpyxl.load_workbook(file_label)
                    ws_label = wb_label.active
                    
                    col_map = {}
                    for c in range(1, ws_label.max_column + 1):
                        val = clean_str(ws_label.cell(row=1, column=c).value)
                        if val: col_map[val] = c
                    col_teacher = col_map.get('任課教師', 6)  
                    col_proctor = col_map.get('監考老師', 8)  
                    
                    def get_val(r, c_idx):
                        v = ws_label.cell(row=r, column=c_idx).value
                        return str(v).strip() if v is not None else ""

                    for r in range(2, ws_label.max_row + 1):
                        # 強制讀取 A, B, D, E 座標
                        val_A = get_val(r, 1) # 序號
                        val_B = ws_label.cell(row=r, column=2).value # 日期
                        val_D = get_val(r, 4) # 班級
                        val_E = get_val(r, 5) # 科目
                        
                        if not val_D: continue
                        
                        cls = clean_str(val_D)
                        subj = normalize_subject(val_E)
                        
                        teacher = get_teacher_fuzzy(cls, subj, course_dict)
                        if teacher:
                            ws_label.cell(row=r, column=col_teacher).value = teacher
                        
                        try: p_val = int(float(val_A))
                        except: p_val = -1
                        
                        if cls in assign_map and 1 <= p_val <= 5:
                            # 讀取標籤日期並轉字串
                            if isinstance(val_B, datetime):
                                date_str = val_B.strftime('%Y-%m-%d')
                            else:
                                date_str = str(val_B).strip().replace('/', '-')
                                
                            day_offset = -1
                            # 【核心修復】：利用檔案真實日期的後五碼(例如 05-14)去配對，無視網頁輸入！
                            if len(sys_d1_raw) >= 5 and sys_d1_raw[-5:] in date_str: day_offset = 0
                            elif len(sys_d2_raw) >= 5 and sys_d2_raw[-5:] in date_str: day_offset = 5
                            
                            if day_offset != -1:
                                target_col = day_offset + p_val
                                proctor = df_assign.iloc[assign_map[cls], target_col]
                                if pd.notna(proctor) and proctor is not None:
                                    ws_label.cell(row=r, column=col_proctor).value = str(proctor).replace('None', '')

                    out_label = io.BytesIO()
                    wb_label.save(out_label)
                    label_bytes = out_label.getvalue()

            st.balloons()
            st.session_state['results'] = {
                'orig': to_excel_bytes(df_out_master, header_df),
                'assign': to_excel_bytes(df_assign, assign_header),
                'pub': pub_bytes,
                'label': label_bytes
            }
            
            st.success(f"🎉 標籤列印已合成！\n\n系統已自動抓取您的真實考試日為：**{sys_d1_raw}** 與 **{sys_d2_raw}**。只要標籤是這兩天的，監考老師保證填得滿滿的！")

        except Exception as e:
            st.error(f"發生錯誤: {e}")
            st.code(traceback.format_exc())

# ==========================================
# 5. 下載區
# ==========================================
if st.session_state['results']:
    st.divider()
    res = st.session_state['results']
    c1, c2, c3, c4 = st.columns(4)
    with c1: st.download_button("📥 1. 監考總表", res['orig'], "監考總表.xlsx", "application/vnd.ms-excel", use_container_width=True)
    with c2: st.download_button("📥 2. 監考一覽表", res['assign'], "監考一覽表_分配完成.xlsx", "application/vnd.ms-excel", use_container_width=True, type="primary")
    with c3: 
        if res['pub']: st.download_button("📥 3. 公布版套印總表", res['pub'], "公布版總表.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
    with c4:
        if res.get('label'): st.download_button("📥 4. 標籤列印(完整)", res['label'], "標籤列印_完整版.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True, type="primary")
